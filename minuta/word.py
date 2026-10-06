"""La bozza in Word, e ritorno.

Minuta consegna ogni bozza anche come file .docx, impaginato come chiede
l'art. 6 del DM 110/2023 (testo vigente letto su Normattiva il 6 ottobre
2026): caratteri di tipo corrente, «preferibilmente» di 12 punti, interlinea
1,5, margini di 2,5 centimetri. Lo studio può cambiare queste impostazioni in
config/studio.json, oppure usare la propria carta intestata.

Nel file l'avvocato trova:
- in cima, un riquadro con le note di Minuta: chi ha preparato la bozza, gli
  avvisi, lo stato delle citazioni;
- a margine, un commento per paragrafo con la provenienza. Commenti, non note
  a piè di pagina: il comma 2 dello stesso articolo non consente note, salvo
  per giurisprudenza e dottrina;
- evidenziati in giallo, i dati da completare.

Riquadro e commenti sono appunti di lavoro: prima del deposito si tolgono.
Quando l'atto firmato torna come .docx, Minuta lo legge come se tutte le
revisioni fossero accettate, salta il proprio riquadro e segnala ciò che
nell'atto non dovrebbe più esserci.
"""

from __future__ import annotations

import re
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .archive import ruolo_del_titolo
from .learning import dividi_note, quanti

IMPAGINAZIONE = {"carattere": "Times New Roman", "corpo": 12, "interlinea": 1.5,
                 "margini_cm": 2.5}
AUTORE = "Minuta"
STILE_NOTA, ID_NOTA = "Minuta Nota", "MinutaNota"
INTESTAZIONE = "BOZZA · fascicolo"
APPUNTI = "Questo riquadro e i commenti a margine sono appunti di Minuta: prima del deposito vanno tolti."

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
# Dentro un paragrafo: ciò che non è testo dell'atto. Il testo cancellato con le
# revisioni (del, moveFrom) non c'è più, una volta accettate; i codici dei campi
# e le caselle di testo non sono il testo dell'atto.
SALTATI = {W + t for t in ("del", "moveFrom", "delText", "instrText", "pPr", "rPr", "drawing",
                            "pict", "commentReference", "footnoteReference", "endnoteReference")}
SALTATI.add("{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent")
CONTENITORI = {W + t for t in ("tbl", "tr", "tc", "sdt", "sdtContent", "customXml")}
# L'ordine delle proprietà di tabella e di cella nello standard OOXML: Word
# rifiuta un file con gli elementi fuori posto.
ORDINE_TABELLA = ("tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize "
                  "tblW jc tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook "
                  "tblCaption tblDescription tblPrChange").split()
ORDINE_CELLA = ("cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection "
                "tcFitText vAlign hideMark headers cellIns cellDel cellMerge tcPrChange").split()
SEGMENTI = re.compile(r"(\*\*[^*\n]+\*\*|\*[^*\n]+\*|\[DA COMPLETARE[^\]]*\])")


# --- Dalla bozza al file Word -------------------------------------------------

def da_markdown(markdown: str) -> dict:
    """La bozza salvata da Minuta (draft.in_markdown), di nuovo divisa nelle sue parti."""
    testa, coda = dividi_note(markdown)
    righe = testa.split("\n")
    intestazione = ""
    if righe and righe[0].startswith(">"):
        intestazione = re.sub(r"[*>]", "", righe[0]).strip()
        righe = righe[1:]
    paragrafi = []
    for blocco in re.split(r"\n\s*\n", "\n".join(righe).strip()):
        blocco = blocco.strip()
        fonte = re.fullmatch(r"<sub>provenienza:\s*([^<]*)</sub>", blocco)
        if fonte and paragrafi:
            paragrafi[-1]["fonte"] = fonte.group(1).strip()
        elif blocco:
            paragrafi.append({"testo": blocco, "fonte": None})
    sezioni, corrente = {}, None
    for riga in coda.split("\n"):
        if riga.startswith("## "):
            corrente = sezioni.setdefault(riga[3:].strip(), [])
        elif riga.startswith("- ") and corrente is not None:
            corrente.append(riga[2:].strip())
    return {"intestazione": intestazione, "paragrafi": paragrafi,
            "note": sezioni.get("Note per l'avvocato", []),
            "citazioni": sezioni.get("Citazioni", [])}


def _documento(impaginazione: dict | None, carta: Path | None):
    if carta is not None and carta.exists():
        # La carta intestata dello studio: intestazione, piè di pagina e stili
        # restano i suoi; del corpo si tiene solo l'impostazione della pagina.
        documento = Document(str(carta))
        corpo = documento.element.body
        for figlio in list(corpo):
            if figlio.tag != qn("w:sectPr"):
                corpo.remove(figlio)
        return documento
    regole = {**IMPAGINAZIONE, **(impaginazione or {})}
    documento = Document()
    sezione = documento.sections[0]
    sezione.page_width, sezione.page_height = Cm(21), Cm(29.7)
    for lato in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sezione, lato, Cm(regole["margini_cm"]))
    normale = documento.styles["Normal"]
    normale.font.name = regole["carattere"]
    normale.font.size = Pt(regole["corpo"])
    formato = normale.paragraph_format
    formato.line_spacing = regole["interlinea"]
    formato.space_after = Pt(6)
    formato.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    lingua = OxmlElement("w:lang")
    lingua.set(qn("w:val"), "it-IT")  # il controllo ortografico di Word, in italiano
    normale.element.get_or_add_rPr().append(lingua)
    return documento


def _metti(proprieta, elemento, ordine: list[str]) -> None:
    nome = elemento.tag.split("}")[1]
    proprieta.insert_element_before(elemento, *(f"w:{n}" for n in ordine[ordine.index(nome) + 1:]))


def _stile_nota(documento):
    stili = documento.styles
    if any(s.name == STILE_NOTA for s in stili):
        return stili[STILE_NOTA]
    stile = stili.add_style(STILE_NOTA, WD_STYLE_TYPE.PARAGRAPH)
    stile.base_style = stili["Normal"]
    stile.font.name = "Arial"
    stile.font.size = Pt(9)
    stile.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
    stile.paragraph_format.line_spacing = 1.0
    stile.paragraph_format.space_after = Pt(2)
    stile.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return stile


def _riquadro(documento, righe: list[str]) -> None:
    """Le note di Minuta in una tabella a una cella, bordata e grigia: si vede
    uguale in ogni programma, e l'avvocato la toglie con un solo gesto."""
    tabella = documento.add_table(rows=1, cols=1)
    bordi = OxmlElement("w:tblBorders")
    for lato in ("top", "left", "bottom", "right"):
        linea = OxmlElement(f"w:{lato}")
        for chiave, valore in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", "999999")):
            linea.set(qn(f"w:{chiave}"), valore)
        bordi.append(linea)
    _metti(tabella._tbl.tblPr, bordi, ORDINE_TABELLA)
    cella = tabella.cell(0, 0)
    sfondo = OxmlElement("w:shd")
    for chiave, valore in (("val", "clear"), ("color", "auto"), ("fill", "F2F2F2")):
        sfondo.set(qn(f"w:{chiave}"), valore)
    _metti(cella._tc.get_or_add_tcPr(), sfondo, ORDINE_CELLA)
    nota = _stile_nota(documento)
    primo = cella.paragraphs[0]
    primo.style = nota
    primo.add_run(righe[0])
    for riga in righe[1:]:
        cella.add_paragraph(riga, style=nota)
    documento.add_paragraph()  # uno spazio fra il riquadro e l'atto


def _maiuscolo(testo: str) -> bool:
    lettere = [c for c in testo if c.isalpha()]
    return len(lettere) >= 3 and all(c.isupper() for c in lettere)


def _riga(documento, riga: str) -> list:
    """Una riga della bozza, come paragrafo di Word: titoli, grassetti, dati da completare."""
    riga = re.sub(r"^#+\s*", "", riga)
    paragrafo = documento.add_paragraph()
    pulita = re.sub(r"[*#]", "", riga).strip()
    titolo = len(pulita) <= 80 and (_maiuscolo(pulita) or ruolo_del_titolo(pulita) is not None)
    if titolo and _maiuscolo(pulita):
        paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pezzi = []
    for pezzo in SEGMENTI.split(riga):
        # Come learning.testo_della_bozza: senza asterischi e cancelletti, così la
        # bozza in Word e quella in Markdown hanno lo stesso testo.
        testo = re.sub(r"[*#]", "", pezzo)
        if not testo:
            continue
        run = paragrafo.add_run(testo)
        if titolo or pezzo.startswith("**"):
            run.bold = True
        elif pezzo.startswith("*") and pezzo.endswith("*") and len(pezzo) > 2:
            run.italic = True
        if testo.startswith("[DA COMPLETARE"):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        pezzi.append(run)
    return pezzi


def scrivi(markdown: str, percorso: Path, impaginazione: dict | None = None,
           carta: Path | None = None) -> Path:
    """La bozza come file Word, pronta da correggere."""
    bozza = da_markdown(markdown)
    documento = _documento(impaginazione, carta)
    riquadro = [bozza["intestazione"], APPUNTI, "Note per l'avvocato:",
                *(f"• {n}" for n in bozza["note"])]
    if bozza["citazioni"]:
        # Una citazione ripetuta nell'atto si controlla una volta: qui basta una riga.
        riquadro += ["Citazioni:", *(f"• {c}" for c in dict.fromkeys(bozza["citazioni"]))]
    _riquadro(documento, riquadro)
    for paragrafo in bozza["paragrafi"]:
        pezzi = []
        for riga in paragrafo["testo"].split("\n"):
            if riga.strip():
                pezzi += _riga(documento, riga.strip())
        if pezzi and paragrafo["fonte"]:
            documento.add_comment(pezzi, text=f"Provenienza: {paragrafo['fonte']}",
                                  author=AUTORE, initials="M")
    percorso.parent.mkdir(parents=True, exist_ok=True)
    documento.save(str(percorso))
    return percorso


# --- Dal file Word all'atto firmato ------------------------------------------

def _xml(percorso: Path, parte: str):
    with zipfile.ZipFile(percorso) as archivio:
        if parte not in archivio.namelist():
            return None
        return ET.fromstring(archivio.read(parte))


def _blocchi(elemento):
    """I paragrafi del corpo nell'ordine di lettura, anche dentro le tabelle."""
    for figlio in elemento:
        if figlio.tag == W + "p":
            yield figlio
        elif figlio.tag in CONTENITORI:
            yield from _blocchi(figlio)


def _testo(elemento) -> str:
    parti = []
    for figlio in elemento:
        if figlio.tag in SALTATI:
            continue
        if figlio.tag == W + "t":
            parti.append(figlio.text or "")
        elif figlio.tag == W + "tab":
            parti.append("\t")
        elif figlio.tag in (W + "br", W + "cr"):
            parti.append("\n")
        elif figlio.tag == W + "noBreakHyphen":
            parti.append("-")
        else:
            parti.append(_testo(figlio))
    return "".join(parti)


def _stile(paragrafo) -> str | None:
    stile = paragrafo.find(f"{W}pPr/{W}pStyle")
    return stile.get(W + "val") if stile is not None else None


def _di_minuta(paragrafo) -> bool:
    testo = _testo(paragrafo)
    return _stile(paragrafo) == ID_NOTA or testo.startswith(INTESTAZIONE) or testo == APPUNTI


def leggi(percorso: Path) -> list[str]:
    """I paragrafi dell'atto, come se tutte le revisioni fossero accettate,
    senza il riquadro di Minuta."""
    corpo = _xml(percorso, "word/document.xml").find(W + "body")
    righe, unisci = [], False
    for paragrafo in _blocchi(corpo):
        if _di_minuta(paragrafo):
            continue
        testo = _testo(paragrafo)
        if unisci and righe:
            righe[-1] += testo
        else:
            righe.append(testo)
        # Un a capo cancellato con le revisioni unisce il paragrafo al successivo.
        unisci = paragrafo.find(f"{W}pPr/{W}rPr/{W}del") is not None
    return [r.strip() for r in righe if r.strip()]


def residui_nel_testo(testo: str) -> list[str]:
    """Ciò che di una bozza di Minuta resta nel testo dell'atto firmato."""
    trovati = []
    if INTESTAZIONE in testo or APPUNTI in testo:
        trovati.append("il riquadro di Minuta")
    if "Provenienza: " in testo:
        trovati.append("i commenti di Minuta sulla provenienza")
    mancanti = testo.count("[DA COMPLETARE")
    if mancanti:
        trovati.append(quanti(mancanti, "dato da completare", "dati da completare"))
    return trovati


def residui(percorso: Path) -> list[str]:
    """Ciò che di Minuta resta nel file Word: in un atto firmato non dovrebbe esserci."""
    documento = _xml(percorso, "word/document.xml")
    trovati = []
    riquadro = sum(1 for p in _blocchi(documento.find(W + "body")) if _di_minuta(p))
    if riquadro:
        trovati.append(f"il riquadro di Minuta ({quanti(riquadro, 'paragrafo', 'paragrafi')})")
    commenti = _xml(percorso, "word/comments.xml")
    if commenti is not None:
        di_minuta = sum(1 for c in commenti.iter(W + "comment") if c.get(W + "author") == AUTORE)
        if di_minuta:
            trovati.append(quanti(di_minuta, "commento di Minuta", "commenti di Minuta"))
    revisioni = sum(1 for nome in ("ins", "del", "moveFrom", "moveTo")
                    for _ in documento.iter(W + nome))
    if revisioni:
        trovati.append(quanti(revisioni, "revisione non accettata", "revisioni non accettate"))
    return trovati + residui_nel_testo("\n".join(leggi(percorso)))
