"""I documenti della lezione B8, inventati per il libro.

    .venv/bin/python basi/genera_documenti.py

Scrive in basi/documenti/ il decreto ingiuntivo di Ristorazione Collinare in
due forme: un PDF con il testo, e la stessa pagina come immagine, come esce da
uno scanner. Poi un promemoria per il cliente in Word, e la fattura n. 112 di
Termocucine Secchia nel formato FatturaPA, convalidata sullo schema ufficiale
e firmata con un certificato di prova creato al momento e non conservato.
Società, partite IVA (con la cifra di controllo sbagliata apposta) e numeri
sono inventati.
"""

import subprocess
import sys
import tempfile
from pathlib import Path

import docx
import pymupdf

QUI = Path(__file__).resolve().parent / "documenti"
sys.path.insert(0, str(QUI.parents[1] / "strumenti"))
import genera_fatture  # noqa: E402

DECRETO = """TRIBUNALE DI MODENA
Decreto ingiuntivo n. 1873/2026

Il giudice, letto il ricorso depositato il 2 settembre 2026 da Termocucine
Secchia S.r.l. contro Ristorazione Collinare S.r.l., ritenuto che il credito
è fondato su prova scritta, visti gli articoli 633 e seguenti del codice di
procedura civile,

INGIUNGE

a Ristorazione Collinare S.r.l. di pagare alla ricorrente, entro quaranta
giorni dalla notificazione di questo decreto, la somma di euro 14.280,00,
oltre agli interessi e alle spese del procedimento, con l'avvertimento che
nello stesso termine può essere fatta opposizione e che, in mancanza di
opposizione, si procederà a esecuzione forzata.

Modena, 15 settembre 2026
Il giudice"""

PROMEMORIA = [
    "Ristorazione Collinare S.r.l. · promemoria sul decreto ingiuntivo",
    "Il decreto n. 1873/2026 è stato notificato il 21 settembre 2026. Il "
    "termine per proporre opposizione è scritto nel decreto e si calcola "
    "dalla notificazione.",
    "Prima di decidere servono la contabile del bonifico di luglio e la PEC "
    "di aprile sul forno.",
]

CEDENTE = {"denominazione": "Termocucine Secchia S.r.l.", "piva": "04811230367",
           "indirizzo": "Via dei Ceramisti", "civico": "12", "cap": "41049",
           "comune": "Sassuolo", "provincia": "MO"}
CESSIONARIO = {"denominazione": "Ristorazione Collinare S.r.l.", "piva": "02931840365",
               "indirizzo": "Via delle Colline", "civico": "4", "cap": "41053",
               "comune": "Maranello", "provincia": "MO"}
FATTURA = {"numero": "112", "data": "2026-02-16", "totale": "6100.00",
           "scadenza": "2026-03-18", "ordine": ("CF-2026-01", "2026-01-20"),
           "descrizione": "Fornitura e installazione di cucina professionale, primo acconto"}


def pdf(percorso: Path) -> None:
    documento = pymupdf.open()
    pagina = documento.new_page()  # A4
    pagina.insert_textbox(pymupdf.Rect(72, 72, 523, 770), DECRETO, fontsize=11)
    documento.set_metadata({"title": "Decreto ingiuntivo n. 1873/2026", "producer": "Minuta, lezione B8"})
    documento.save(percorso, deflate=True, garbage=3)


def scansione(da: Path, percorso: Path) -> None:
    """La stessa pagina come immagine: nessun testo da estrarre."""
    with pymupdf.open(da) as originale:
        # In toni di grigio e compressa, come da uno scanner da ufficio.
        immagine = originale[0].get_pixmap(dpi=100, colorspace=pymupdf.csGRAY)
        documento = pymupdf.open()
        pagina = documento.new_page(width=originale[0].rect.width, height=originale[0].rect.height)
        pagina.insert_image(pagina.rect, stream=immagine.tobytes("png"))
        documento.set_metadata({"title": "Scansione", "producer": "Minuta, lezione B8"})
        documento.save(percorso, deflate=True, garbage=3)


def word(percorso: Path) -> None:
    documento = docx.Document()
    documento.add_heading(PROMEMORIA[0], level=1)
    for paragrafo in PROMEMORIA[1:]:
        documento.add_paragraph(paragrafo)
    documento.core_properties.author = "Studio Meridiana"
    documento.save(percorso)


def fattura_firmata(percorso: Path) -> None:
    xml = genera_fatture.fattura(CEDENTE, CESSIONARIO, [FATTURA], "00112")
    genera_fatture.valida(xml)
    with tempfile.TemporaryDirectory() as cartella:
        c = Path(cartella)
        (c / "fattura.xml").write_bytes(xml)
        subprocess.run(["openssl", "req", "-x509", "-newkey", "rsa:2048", "-nodes",
                        "-keyout", str(c / "chiave.pem"), "-out", str(c / "cert.pem"),
                        "-days", "3650", "-subj", "/CN=Termocucine Secchia (certificato di prova)"],
                       check=True, capture_output=True)
        subprocess.run(["openssl", "smime", "-sign", "-binary", "-nodetach",
                        "-in", str(c / "fattura.xml"), "-signer", str(c / "cert.pem"),
                        "-inkey", str(c / "chiave.pem"), "-outform", "DER",
                        "-out", str(percorso)], check=True, capture_output=True)


if __name__ == "__main__":
    pdf(QUI / "decreto.pdf")
    scansione(QUI / "decreto.pdf", QUI / "decreto-scansione.pdf")
    word(QUI / "promemoria.docx")
    fattura_firmata(QUI / "fattura.xml.p7m")
    for nome in ("decreto.pdf", "decreto-scansione.pdf", "promemoria.docx", "fattura.xml.p7m"):
        print(nome, (QUI / nome).stat().st_size, "byte")
