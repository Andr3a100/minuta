"""I documenti dei tre fascicoli di prova, inventati per il libro (lezione 16).

    .venv/bin/python strumenti/genera_documenti_di_prova.py

Scrive in documenti-di-prova/ una cartella per fascicolo:
- 2026-071, Ristorazione Collinare (civile): il ricorso per decreto
  ingiuntivo della lezione 11 e il decreto, come PDF con il testo; le tre
  fatture di Termocucine Secchia nel formato FatturaPA, convalidate sullo
  schema ufficiale e firmate con un certificato di prova; il promemoria per
  il cliente in Word;
- 2026-072, Ivo Marchetti (tributario): la cartella di pagamento, come la
  scansiona uno scanner da ufficio, a 100 e a 300 punti per pollice;
- 2026-073, Alessandro Riva (penale): l'avviso di conclusione delle indagini,
  un PDF con il testo dentro una busta firmata con un certificato di prova.

I certificati di prova si creano al momento e non si conservano. Persone,
società, partite IVA, numeri e fatti sono inventati.
"""

from __future__ import annotations

import subprocess
import sys
import tempfile
from pathlib import Path

import docx
import pymupdf

RADICE = Path(__file__).resolve().parents[1]
QUI = RADICE / "documenti-di-prova"
sys.path.insert(0, str(RADICE / "strumenti"))
import genera_fatture  # noqa: E402

DECRETO = """TRIBUNALE DI MODENA
Decreto ingiuntivo n. 1873/2026

Il giudice, letto il ricorso depositato il 2 settembre 2026 da Termocucine
Secchia S.r.l. contro Ristorazione Collinare S.r.l., ritenuto che il credito
è fondato su prova scritta, visti gli articoli 633 e seguenti del codice di
procedura civile,

INGIUNGE

a Ristorazione Collinare S.r.l. di pagare alla ricorrente, entro quaranta
giorni dalla notificazione di questo decreto, la somma di euro 14.280,00,
oltre agli interessi e alle spese del procedimento, con l'avvertimento che
nello stesso termine può essere fatta opposizione e che, in mancanza di
opposizione, si procederà a esecuzione forzata.

Modena, 15 settembre 2026
Il giudice"""

PROMEMORIA = [
    "Ristorazione Collinare S.r.l. · promemoria sul decreto ingiuntivo",
    "Il decreto n. 1873/2026 è stato notificato il 21 settembre 2026. Il "
    "termine per proporre opposizione è scritto nel decreto e si calcola "
    "dalla notificazione.",
    "Prima di decidere servono la contabile del bonifico di luglio e la PEC "
    "di aprile sul forno.",
]

CEDENTE = {
    "denominazione": "Termocucine Secchia S.r.l.",
    "piva": "04811230367",
    "indirizzo": "Via dei Ceramisti",
    "civico": "12",
    "cap": "41049",
    "comune": "Sassuolo",
    "provincia": "MO",
}
CESSIONARIO = {
    "denominazione": "Ristorazione Collinare S.r.l.",
    "piva": "02931840365",
    "indirizzo": "Via delle Colline",
    "civico": "4",
    "cap": "41053",
    "comune": "Maranello",
    "provincia": "MO",
}
CUCINA = "Fornitura e installazione di cucina professionale"
# Le tre fatture del ricorso della lezione 11: 14.280,00 euro in tutto.
FATTURE = [
    {
        "numero": "112",
        "data": "2026-02-16",
        "totale": "6100.00",
        "scadenza": "2026-03-18",
        "ordine": ("CF-2026-01", "2026-01-20"),
        "descrizione": f"{CUCINA}, primo acconto",
    },
    {
        "numero": "141",
        "data": "2026-03-10",
        "totale": "5480.00",
        "scadenza": "2026-04-09",
        "ordine": ("CF-2026-01", "2026-01-20"),
        "ddt": ("27", "2026-03-04"),
        "descrizione": f"{CUCINA}, secondo acconto",
    },
    {
        "numero": "168",
        "data": "2026-03-31",
        "totale": "2700.00",
        "scadenza": "2026-04-30",
        "ordine": ("CF-2026-01", "2026-01-20"),
        "ddt": ("31", "2026-03-09"),
        "descrizione": f"{CUCINA}, saldo",
    },
]

# La cartella di Ivo Marchetti: le cifre in colonna, come nei moduli.
CARTELLA = """CARTELLA DI PAGAMENTO N. 070 2026 00418265 31 000

Contribuente: MARCHETTI IVO
Domicilio fiscale: Via dei Ciliegi 14, 41058 Vignola (MO)

Ente creditore: Agenzia delle Entrate, Direzione provinciale di Modena
Ruolo n. 2026/004871, reso esecutivo il 14 luglio 2026

Anno d'imposta 2022 - IRPEF - controllo automatizzato della dichiarazione

Imposta                          3.412,00
Sanzioni                         1.023,60
Interessi                          396,18
Diritti di notifica                  5,88
Totale da pagare                 4.837,66

Si intima di pagare entro sessanta giorni dalla notificazione della
presente cartella; in mancanza si procederà a esecuzione forzata. Entro lo
stesso termine può essere proposto ricorso alla Corte di giustizia
tributaria di primo grado.

RELATA DI NOTIFICA
Notificata il 18 settembre 2026 mediante consegna a mani del destinatario."""

# I fatti sono quelli del caso della lezione 4. I due avvertimenti
# dell'avviso riprendono i commi 2 e 3 dell'art. 415-bis del codice di
# procedura penale, testo vigente al 7 ottobre 2026.
AVVISO = """PROCURA DELLA REPUBBLICA PRESSO IL TRIBUNALE DI MODENA
Procedimento penale n. 3127/2026 R.G.N.R.

AVVISO ALL'INDAGATO DELLA CONCLUSIONE DELLE INDAGINI PRELIMINARI
(art. 415-bis del codice di procedura penale)

Il pubblico ministero, visti gli atti del procedimento,

AVVISA

Alessandro Riva, legale rappresentante di Logistica Riva S.r.l., indagato,
e il suo difensore, avv. Stefano Valli del foro di Modena, che le indagini
preliminari sono concluse per il seguente fatto: lesioni personali colpose
gravi, commesse con violazione delle norme per la prevenzione degli
infortuni sul lavoro (art. 590, secondo e terzo comma, del codice penale),
in danno del dipendente Marco Bellini, caduto da un soppalco nel magazzino
della società a Carpi il 3 marzo 2026, con una frattura del bacino e una
prognosi di novanta giorni.

AVVERTE

che la documentazione relativa alle indagini espletate è depositata presso
la segreteria del pubblico ministero e che l'indagato e il suo difensore
hanno facoltà di prenderne visione ed estrarne copia;

che l'indagato ha facoltà, entro il termine di venti giorni, di presentare
memorie, produrre documenti, depositare documentazione relativa ad
investigazioni del difensore, chiedere al pubblico ministero il compimento
di atti di indagine, nonché di presentarsi per rilasciare dichiarazioni
ovvero chiedere di essere sottoposto ad interrogatorio;

che l'indagato e la persona offesa hanno facoltà di accedere ai programmi
di giustizia riparativa.

Modena, 1° ottobre 2026
Il pubblico ministero"""


def pdf_con_testo(percorso: Path, testo: str, titolo: str, corpo=11) -> None:
    documento = pymupdf.open()
    pagina = documento.new_page()  # A4
    avanzo = pagina.insert_textbox(
        pymupdf.Rect(60, 60, 535, 790), testo, fontsize=corpo
    )
    assert avanzo >= 0, f"{percorso.name}: il testo non sta nella pagina"
    documento.set_metadata({"title": titolo, "producer": "Minuta, lezione 16"})
    documento.save(percorso, deflate=True, garbage=3)


def scansione(percorso: Path, testo: str, punti: int) -> None:
    """La pagina come esce da uno scanner: un'immagine, nessun testo."""
    originale = pymupdf.open()
    pagina = originale.new_page()
    avanzo = pagina.insert_textbox(
        pymupdf.Rect(60, 60, 540, 800), testo, fontname="cour", fontsize=10
    )
    assert avanzo >= 0, "la cartella non sta nella pagina"
    # In toni di grigio, come da uno scanner da ufficio.
    immagine = pagina.get_pixmap(dpi=punti, colorspace=pymupdf.csGRAY)
    documento = pymupdf.open()
    nuova = documento.new_page(
        width=pagina.rect.width, height=pagina.rect.height
    )
    nuova.insert_image(nuova.rect, stream=immagine.tobytes("png"))
    documento.set_metadata({"title": "Scansione", "producer": "Minuta"})
    documento.save(percorso, deflate=True, garbage=3)


def certificato(cartella: Path, nome: str) -> tuple[Path, Path]:
    chiave, cert = cartella / "chiave.pem", cartella / "cert.pem"
    subprocess.run(
        [
            "openssl",
            "req",
            "-x509",
            "-newkey",
            "rsa:2048",
            "-nodes",
            "-keyout",
            str(chiave),
            "-out",
            str(cert),
            "-days",
            "3650",
            "-subj",
            f"/CN={nome}",
        ],
        check=True,
        capture_output=True,
    )
    return chiave, cert


def firma(dati: bytes, percorso: Path, chiave: Path, cert: Path) -> None:
    subprocess.run(
        [
            "openssl",
            "smime",
            "-sign",
            "-binary",
            "-nodetach",
            "-signer",
            str(cert),
            "-inkey",
            str(chiave),
            "-outform",
            "DER",
            "-out",
            str(percorso),
        ],
        input=dati,
        check=True,
        capture_output=True,
    )


def civile(cartella: Path) -> None:
    ricorso = (RADICE / "esempi" / "11-ricorso.txt").read_text("utf-8")
    pdf_con_testo(cartella / "ricorso.pdf", ricorso, "Ricorso", corpo=9.5)
    pdf_con_testo(cartella / "decreto.pdf", DECRETO, "Decreto ingiuntivo")
    with tempfile.TemporaryDirectory() as temporanea:
        chiave, cert = certificato(
            Path(temporanea), "Termocucine Secchia (certificato di prova)"
        )
        for numero, fattura in enumerate(FATTURE, start=1):
            xml = genera_fatture.fattura(
                CEDENTE, CESSIONARIO, [fattura], f"{numero:05d}"
            )
            genera_fatture.valida(xml)
            nome = f"fattura-{fattura['numero']}.xml.p7m"
            firma(xml, cartella / nome, chiave, cert)
    promemoria = docx.Document()
    promemoria.add_heading(PROMEMORIA[0], level=1)
    for paragrafo in PROMEMORIA[1:]:
        promemoria.add_paragraph(paragrafo)
    promemoria.core_properties.author = "Studio Meridiana"
    promemoria.save(cartella / "promemoria.docx")


def tributario(cartella: Path) -> None:
    scansione(cartella / "cartella-scansione.pdf", CARTELLA, 100)
    scansione(cartella / "cartella-scansione-300.pdf", CARTELLA, 300)


def penale(cartella: Path) -> None:
    with tempfile.TemporaryDirectory() as temporanea:
        avviso = Path(temporanea) / "avviso.pdf"
        pdf_con_testo(avviso, AVVISO, "Avviso art. 415-bis", corpo=10.5)
        chiave, cert = certificato(
            Path(temporanea),
            "Procura della Repubblica di Modena (certificato di prova)",
        )
        firma(
            avviso.read_bytes(),
            cartella / "avviso-415-bis.pdf.p7m",
            chiave,
            cert,
        )


if __name__ == "__main__":
    for codice, prepara in (
        ("2026-071", civile),
        ("2026-072", tributario),
        ("2026-073", penale),
    ):
        cartella = QUI / codice
        cartella.mkdir(parents=True, exist_ok=True)
        prepara(cartella)
        for file in sorted(cartella.iterdir()):
            print(f"{codice}/{file.name}", file.stat().st_size, "byte")
