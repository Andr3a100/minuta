"""Le correzioni dell'avvocato, nell'applicazione (lezione 26)."""

from __future__ import annotations

import importlib.util
import json
from argparse import Namespace
from pathlib import Path

import pytest
from aiuti import invia
from docx import Document
from sqlalchemy import func, inspect, select
from test_app_bozze import ModelloCheRicorda, civile, utente  # noqa: F401

from app import documenti as app_documenti
from app import manage
from app.bozze import BozzaRifiutata, approva_bozza, in_word, prepara_bozza
from app.correzioni import (
    CorrezioneRifiutata,
    regole_di,
    respingi,
    stato_di,
    voci_di,
)
from app.db import (
    Bozza,
    Correzione,
    Evento,
    Fascicolo,
    PersonaFascicolo,
    make_engine,
)
from app.invio import aggiungi_persona
from app.migrations import PASSI, migra
from app.stile import calcola_profilo, letto, ultimo_profilo
from minuta.model import ModelloFinto

RADICE = Path(__file__).resolve().parents[2]
SECONDO = RADICE / "documenti-di-prova" / "2026-074"
_spec = importlib.util.spec_from_file_location(
    "correggi_in_word", RADICE / "percorso" / "correggi_in_word.py"
)
correggi_in_word = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(correggi_in_word)


def secondo_fascicolo(db, impostazioni, sarti) -> Fascicolo:
    """Lavanderia Sassolese, il secondo fascicolo civile di Elena Sarti,
    aperto dal terminale come nella lezione."""
    manage.apri_dalla_riga(
        db,
        Namespace(
            codice="2026-074",
            cliente="Lavanderia Sassolese S.r.l.",
            materia="civile",
            oggetto="Decreto ingiuntivo chiesto da Scoltenna Forniture",
            avvocato="sarti",
            praticante=None,
            da="rosa",
        ),
    )
    fascicolo = db.scalar(
        select(Fascicolo).where(Fascicolo.codice == "2026-074")
    )
    app_documenti.carica(
        db,
        impostazioni,
        sarti,
        fascicolo,
        "decreto.pdf",
        (SECONDO / "decreto.pdf").read_bytes(),
    )
    aggiungi_persona(
        db,
        sarti,
        fascicolo,
        "Scoltenna Forniture S.r.l.",
        "controparte",
        soggetto=True,
    )
    return fascicolo


def approva_corretta(db, avvocato, fascicolo, cartella, tipo="opposizione"):
    """La bozza del modello finto, corretta come la corregge Elena Sarti
    nella lezione, e approvata dal file Word."""
    bozza = prepara_bozza(db, avvocato, fascicolo, tipo, ModelloFinto()).bozza
    file = in_word(db, bozza, cartella / f"bozza-{bozza.id}.docx")
    approvato = cartella / f"approvato-{bozza.id}.docx"
    correggi_in_word.main(str(file), str(approvato))
    approva_bozza(db, avvocato, fascicolo, bozza, approvato)
    return bozza


def test_all_approvazione_si_contano_le_modifiche(app, civile, tmp_path):
    with app.state.session_factory() as db:
        sarti, fascicolo = utente(db, "sarti"), db.get(Fascicolo, civile.id)
        bozza = approva_corretta(db, sarti, fascicolo, tmp_path)
        voci = voci_di(db, sarti)
        evento = db.scalar(
            select(Evento).where(Evento.azione == "approvazione")
        )
    assert json.loads(bozza.modifiche) == {
        "completamento": 1,
        "dati": 0,
        "stile": 2,
        "riscrittura": 2,
    }
    assert json.loads(evento.dettaglio)["modifiche"]["stile"] == 2
    # i dati della formula riempita sono completamenti, non stile
    assert [(v["prima"], v["dopo"], stato_di(v)) for v in voci] == [
        ("ricorrente", "opponente", "osservata"),
        ("somma", "sorte capitale", "osservata"),  # un atto solo: un caso?
    ]


# [libro:prova-regola]
def test_in_due_fascicoli_diventa_regola(app, civile, impostazioni, tmp_path):
    modello = ModelloCheRicorda()
    with app.state.session_factory() as db:
        sarti = utente(db, "sarti")
        approva_corretta(db, sarti, db.get(Fascicolo, civile.id), tmp_path)
        assert regole_di(db, sarti) == []  # in un atto solo: osservate
        secondo = secondo_fascicolo(db, impostazioni, sarti)
        approva_corretta(db, sarti, secondo, tmp_path)
        respingi(db, sarti, 1)  # «ricorrente»: dipende dal tipo di atto
        (regola,) = regole_di(db, sarti)
        esito = prepara_bozza(db, sarti, secondo, "opposizione", modello)
    assert regola["fascicoli"] == ["2026-071", "2026-074"]
    assert (
        "REGOLE APPRESE DALLE CORREZIONI:\n- Scrivi «sorte capitale», non "
        "«somma» (corretto in 2 atti).\n\n" in modello.ricevuti[-1]
    )
    segnalate = [
        q for p in esito.paragrafi for q in p.problemi if "regola" in q
    ]
    assert segnalate == [
        "regola appresa non seguita: scrive «somma», che l'avvocato ha "
        "corretto in «sorte capitale» in 2 atti"
    ]


# [/libro:prova-regola]


def test_due_bozze_dello_stesso_fascicolo_non_fanno_una_regola(
    app, civile, tmp_path
):
    with app.state.session_factory() as db:
        sarti, fascicolo = utente(db, "sarti"), db.get(Fascicolo, civile.id)
        approva_corretta(db, sarti, fascicolo, tmp_path)
        approva_corretta(db, sarti, fascicolo, tmp_path)
        voci = voci_di(db, sarti)
        assert regole_di(db, sarti) == []
    assert [(v["volte"], v["fascicoli"]) for v in voci] == [
        (2, ["2026-071"]),
        (2, ["2026-071"]),
    ]


def test_prima_le_correzioni_fatte_in_piu_fascicoli(app, civile, tmp_path):
    """L'ordine dà i numeri che respingi-correzione usa."""
    with app.state.session_factory() as db:
        sarti, fascicolo = utente(db, "sarti"), db.get(Fascicolo, civile.id)
        bozza = approva_corretta(db, sarti, fascicolo, tmp_path)
        altro = db.scalar(
            select(Fascicolo).where(Fascicolo.codice != "2026-071")
        )
        for codice, prima in (
            (fascicolo.id, "ai sensi"),
            (fascicolo.id, "ai sensi"),
            (fascicolo.id, "zeta"),
            (altro.id, "zeta"),
        ):
            db.add(
                Correzione(
                    avvocato_id=sarti.id,
                    bozza_id=bozza.id,
                    fascicolo_id=codice,
                    prima=prima,
                    dopo="x",
                    contesto="",
                )
            )
        db.commit()
        ordine = [v["prima"] for v in voci_di(db, sarti)]
    # zeta in due fascicoli; ai sensi due volte in uno; le altre una volta
    assert ordine == ["zeta", "ai sensi", "ricorrente", "somma"]


def test_una_correzione_respinta_non_diventa_regola(
    app, civile, impostazioni, tmp_path
):
    modello = ModelloCheRicorda()
    with app.state.session_factory() as db:
        sarti = utente(db, "sarti")
        approva_corretta(db, sarti, db.get(Fascicolo, civile.id), tmp_path)
        secondo = secondo_fascicolo(db, impostazioni, sarti)
        with pytest.raises(CorrezioneRifiutata, match="l'avvocato"):
            respingi(db, utente(db, "irene"), 1)
        with pytest.raises(CorrezioneRifiutata, match="non c'è la n. 3"):
            respingi(db, sarti, 3)
        respingi(db, sarti, 2)  # «somma» → «sorte capitale»
        with pytest.raises(CorrezioneRifiutata, match="già respinta"):
            respingi(db, sarti, 2)
        evento = db.scalar(
            select(Evento).where(Evento.azione == "correzione_respinta")
        )
        assert json.loads(evento.dettaglio) == {
            "dopo": "sorte capitale",
            "prima": "somma",
        }
        approva_corretta(db, sarti, secondo, tmp_path)
        respinta = voci_di(db, sarti)[1]
        (regola,) = regole_di(db, sarti)
        esito = prepara_bozza(db, sarti, secondo, "opposizione", modello)
    assert len(respinta["fascicoli"]) == 2
    assert stato_di(respinta).startswith("respinta il ")
    assert regola["prima"] == "ricorrente"
    assert "non «somma»" not in modello.ricevuti[-1]
    assert not any("«somma»" in str(p.problemi) for p in esito.paragrafi)


def test_un_segnaposto_rimasto_non_entra_nell_atto(app, civile, tmp_path):
    with app.state.session_factory() as db:
        sarti, fascicolo = utente(db, "sarti"), db.get(Fascicolo, civile.id)
        bozza = prepara_bozza(
            db, sarti, fascicolo, "opposizione", ModelloFinto()
        ).bozza
        corretto = Document()
        corretto.add_paragraph(
            "Rappresentata dall'Avv. [PERSONA] (C.F. [CF])."
        )
        corretto.add_paragraph("[DA VERIFICARE] il giudice competente.")
        corretto.save(str(tmp_path / "corretto.docx"))
        with pytest.raises(BozzaRifiutata) as rifiuto:
            approva_bozza(
                db, sarti, fascicolo, bozza, tmp_path / "corretto.docx"
            )
    assert (
        "segnaposto rimasto: [PERSONA]; segnaposto rimasto: [CF]; "
        "1 dubbio da verificare" in str(rifiuto.value)
    )


def test_il_promemoria_non_insegna_lo_stile_degli_atti(app, civile, tmp_path):
    with app.state.session_factory() as db:
        sarti, fascicolo = utente(db, "sarti"), db.get(Fascicolo, civile.id)
        bozza = prepara_bozza(
            db, sarti, fascicolo, "promemoria", ModelloFinto()
        ).bozza
        corretto = Document()
        corretto.add_paragraph("Gentile signora, ecco che cosa è arrivato.")
        corretto.save(str(tmp_path / "corretto.docx"))
        approva_bozza(db, sarti, fascicolo, bozza, tmp_path / "corretto.docx")
        correzioni = db.scalar(select(func.count()).select_from(Correzione))
    assert correzioni == 0
    assert sum(json.loads(bozza.modifiche).values()) > 0


def test_gli_atti_approvati_non_entrano_nel_profilo(
    app, civile, impostazioni, tmp_path
):
    with app.state.session_factory() as db:
        sarti = utente(db, "sarti")
        approva_corretta(db, sarti, db.get(Fascicolo, civile.id), tmp_path)
        approva_corretta(
            db, sarti, secondo_fascicolo(db, impostazioni, sarti), tmp_path
        )
        confermato = letto(ultimo_profilo(db, sarti))
        dopo = letto(calcola_profilo(db, sarti))
    assert dopo["atti"] == confermato["atti"]
    assert [f["testo"] for f in dopo["formule"]] == [
        f["testo"] for f in confermato["formule"]
    ]


def test_i_comandi_delle_correzioni(
    app, civile, impostazioni, tmp_path, capsys
):
    with app.state.session_factory() as db:
        sarti = utente(db, "sarti")
        approva_corretta(db, sarti, db.get(Fascicolo, civile.id), tmp_path)
        secondo = secondo_fascicolo(db, impostazioni, sarti)
        bozza = prepara_bozza(
            db, sarti, secondo, "opposizione", ModelloFinto()
        ).bozza
        file = in_word(db, bozza, tmp_path / "bozza.docx")
        correggi_in_word.main(str(file), str(tmp_path / "approvato.docx"))
        capsys.readouterr()
        manage.approva_dalla_riga(
            db,
            Namespace(
                codice="2026-074",
                numero=bozza.id,
                file=str(tmp_path / "approvato.docx"),
                forza=False,
                da="sarti",
            ),
        )
        approvata = capsys.readouterr().out
        manage.correzioni_dalla_riga(db, Namespace(da="sarti"))
        elenco = capsys.readouterr().out
        with pytest.raises(manage.CommandError, match="di un avvocato"):
            manage.correzioni_dalla_riga(db, Namespace(da="irene"))
    assert approvata.splitlines()[1:] == [
        "Modifiche dell'avvocato: 2 correzioni di stile, 1 completamento, "
        "2 riscritture.",
        "Correzioni di stile:",
        "  «ricorrente» → «opponente»: in 2 atti, regola",
        "  «somma» → «sorte capitale»: in 2 atti, regola",
    ]
    assert elenco.splitlines() == [
        "Le correzioni di Elena Sarti: 2.",
        " 1. «ricorrente» → «opponente»",
        "    in 2 atti (2026-071, 2026-074): regola",
        " 2. «somma» → «sorte capitale»",
        "    in 2 atti (2026-071, 2026-074): regola",
        "Nella richiesta delle sue bozze, d'ora in poi:",
        "- Scrivi «opponente», non «ricorrente» (corretto in 2 atti).",
        "- Scrivi «sorte capitale», non «somma» (corretto in 2 atti).",
    ]


def test_la_pagina_delle_correzioni(app, browser, civile, tmp_path):
    with app.state.session_factory() as db:
        sarti = utente(db, "sarti")
        approva_corretta(db, sarti, db.get(Fascicolo, civile.id), tmp_path)
    pagina = browser("sarti").get("/stile").text
    assert "Le mie correzioni" in pagina
    assert "«somma» → «sorte capitale»" in pagina
    assert "in 1 atto (2026-071) · osservata" in pagina
    assert "Le mie correzioni" not in browser("irene").get("/stile").text
    respinta = invia(
        browser("sarti"), "/stile", {"azione": "respingi", "numeri": "1"}
    )
    assert respinta.status_code == 303
    assert "respinta il " in browser("sarti").get("/stile").text
    di_nuovo = invia(
        browser("sarti"), "/stile", {"azione": "respingi", "numeri": "1"}
    )
    assert di_nuovo.status_code == 409


def test_aprire_un_fascicolo_dal_terminale(app, capsys):
    argomenti = dict(
        codice="2026-090",
        cliente="Cliente di prova S.r.l.",
        materia="civile",
        oggetto="Prova",
        avvocato="dini",
        praticante="irene",
    )
    with app.state.session_factory() as db:
        with pytest.raises(manage.CommandError, match="o la segreteria"):
            manage.apri_dalla_riga(db, Namespace(**argomenti, da="irene"))
        manage.apri_dalla_riga(db, Namespace(**argomenti, da="rosa"))
        with pytest.raises(manage.CommandError, match="Esiste già"):
            manage.apri_dalla_riga(db, Namespace(**argomenti, da="dini"))
        sbagliato = {**argomenti, "codice": "90", "materia": "lavoro"}
        with pytest.raises(manage.CommandError, match="anno-numero"):
            manage.apri_dalla_riga(db, Namespace(**sbagliato, da="rosa"))
        fascicolo = db.scalar(
            select(Fascicolo).where(Fascicolo.codice == "2026-090")
        )
        evento = db.scalar(select(Evento).where(Evento.azione == "apertura"))
        rosa = utente(db, "rosa")
        elenco = [
            p.nome
            for p in db.scalars(
                select(PersonaFascicolo).where(
                    PersonaFascicolo.fascicolo_id == fascicolo.id
                )
            )
        ]
    assert capsys.readouterr().out == (
        "Aperto il fascicolo 2026-090: Cliente di prova S.r.l., civile, "
        "avvocato Marco Dini.\n"
    )
    assert len(json.loads(evento.dettaglio)["persone"]) == 3  # con irene
    assert fascicolo.creato_da_id == evento.utente_id == rosa.id
    assert elenco == ["Cliente di prova S.r.l."]  # il cliente, in elenco


def test_la_migrazione_11_aggiunge_le_correzioni(impostazioni):
    tutti = dict(PASSI)
    engine = make_engine(impostazioni.database_url)
    try:
        PASSI.clear()
        PASSI.update({n: tutti[n] for n in range(1, 11)})  # la versione 1.0
        migra(engine)
        PASSI[11] = tutti[11]  # arriva la versione 1.1
        assert migra(engine) == [11]
        with engine.connect() as conn:
            controllo = inspect(conn)
            assert controllo.has_table("correzioni")
            assert controllo.has_table("correzioni_respinte")
            colonne = {c["name"] for c in controllo.get_columns("bozze")}
        assert "modifiche" in colonne
        assert Bozza.__table__.c.modifiche.nullable
    finally:
        PASSI.clear()
        PASSI.update(tutti)
        engine.dispose()
