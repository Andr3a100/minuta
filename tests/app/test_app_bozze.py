"""Le bozze, nell'applicazione (lezione 24)."""

from __future__ import annotations

import json
import re
import zipfile
from argparse import Namespace
from datetime import date
from pathlib import Path

import pytest
from aiuti import con_fascicoli, invia
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Cm, Pt
from sqlalchemy import func, inspect, select

from app import bozze as app_bozze
from app import documenti as app_documenti
from app import manage
from app.archivio import carica_atto
from app.bozze import (
    BozzaRifiutata,
    approva_bozza,
    citazioni_della_bozza,
    in_word,
    norma_letta,
    prepara_bozza,
)
from app.db import (
    Bozza,
    Documento,
    Evento,
    Fascicolo,
    Scadenza,
    Scheda,
    UsoAI,
    Utente,
    make_engine,
)
from app.informativa import registra_consegna
from app.invio import InvioBloccato, aggiungi_persona
from app.migrations import PASSI, migra
from app.registro import RichiestaNonAmmessa
from app.stile import calcola_profilo, conferma_profilo, decidi_formule, letto
from minuta.archivio import leggi_indice
from minuta.bozza import TipoDiBozza
from minuta.model import ModelloFinto, Risposta

RADICE = Path(__file__).resolve().parents[2]
CIVILE = RADICE / "documenti-di-prova" / "2026-071"
ARCHIVIO = RADICE / "archivio"
NOMI_VERI = ("Ristorazione Collinare", "Termocucine Secchia", "Elena Sarti")


class ModelloCheRicorda(ModelloFinto):
    def __init__(self):
        super().__init__()
        self.ricevuti: list[str] = []

    def scrivi(self, sistema, richiesta) -> Risposta:
        self.ricevuti.append(richiesta)
        return super().scrivi(sistema, richiesta)


def utente(db, nome_utente: str) -> Utente:
    return db.scalar(select(Utente).where(Utente.nome_utente == nome_utente))


def con_scadenza(db, fascicolo, sarti) -> None:
    """Una scadenza confermata, come quella della lezione 19."""
    decreto = db.scalar(
        select(Documento).where(
            Documento.fascicolo_id == fascicolo.id,
            Documento.nome == "decreto.pdf",
        )
    )
    voce = UsoAI(
        utente_id=sarti.id,
        fascicolo_id=fascicolo.id,
        modello="prova",
        categorie="",
        scopo="prova",
        impronta="0" * 64,
        caratteri=0,
        controllo="",
    )
    db.add(voce)
    db.flush()
    scheda = Scheda(
        fascicolo_id=fascicolo.id,
        documento_id=decreto.id,
        righe="[]",
        uso_ai_id=voce.id,
        preparata_da_id=sarti.id,
        confermata_da_id=sarti.id,
    )
    db.add(scheda)
    db.flush()
    db.add(
        Scadenza(
            fascicolo_id=fascicolo.id,
            documento_id=decreto.id,
            scheda_id=scheda.id,
            termine="opposizione a decreto ingiuntivo",
            norma="art. 641 c.p.c.",
            partenza=date(2026, 9, 21),
            scadenza=date(2026, 11, 2),
            passaggi="[]",
            calcolata_da_id=sarti.id,
            confermata_da_id=sarti.id,
        )
    )
    db.commit()


@pytest.fixture
def civile(app, impostazioni):
    """Ristorazione Collinare: i documenti, la controparte in elenco, lo
    stile confermato di Elena Sarti e la scadenza confermata."""
    fascicolo = con_fascicoli(app)["2026-071"]
    indice = leggi_indice(ARCHIVIO / "indice.csv")
    with app.state.session_factory() as db:
        sarti = utente(db, "sarti")
        dentro = db.get(Fascicolo, fascicolo.id)
        for documento in sorted(CIVILE.iterdir()):
            app_documenti.carica(
                db,
                impostazioni,
                sarti,
                dentro,
                documento.name,
                documento.read_bytes(),
            )
        aggiungi_persona(
            db,
            sarti,
            dentro,
            "Termocucine Secchia S.r.l.",
            "controparte",
            soggetto=True,
        )
        for pdf in sorted((ARCHIVIO / "pdf").glob("*.pdf")):
            carica_atto(db, impostazioni, sarti, pdf, indice[pdf.stem])
        dati = letto(calcola_profilo(db, sarti))
        aperte = [
            n
            for n, f in enumerate(dati["formule"], start=1)
            if f["stato"] == "da decidere"
        ]
        decidi_formule(db, sarti, aperte, "tolta")
        conferma_profilo(db, sarti)
        con_scadenza(db, dentro, sarti)
        # Il cliente ha ricevuto l'informativa sull'IA (lezione 34).
        registra_consegna(db, sarti, dentro, date(2026, 10, 1), "a mano")
    return fascicolo


# [libro:prova-bozza-app]
def test_la_bozza_parte_senza_nomi_e_torna_controllata(app, civile):
    modello = ModelloCheRicorda()
    with app.state.session_factory() as db:
        esito = prepara_bozza(
            db, utente(db, "sarti"), civile, "opposizione", modello
        )
    (inviato,) = modello.ricevuti
    assert not any(nome in inviato for nome in NOMI_VERI)
    assert "STILE:" in inviato and "SCADENZE:" in inviato
    (pagamento,) = [p for p in esito.paragrafi if "bonifico" in p.testo]
    assert pagamento.problemi == [
        "data che non è nei documenti: mercoledì 15 luglio 2026",
        "importo che non è nei documenti: 5.000,00",
    ]
    assert "Da completare: i motivi dell'opposizione e le loro prove" in (
        esito.note
    )


def test_l_atto_solo_con_lo_stile_confermato(app, impostazioni):
    fascicolo = con_fascicoli(app)["2026-071"]
    with app.state.session_factory() as db:
        with pytest.raises(BozzaRifiutata, match="profilo, confermato"):
            prepara_bozza(
                db,
                utente(db, "irene"),
                fascicolo,
                "opposizione",
                ModelloFinto(),
            )


# [/libro:prova-bozza-app]


class CitaIlCliente(ModelloFinto):
    """Risponde con il primo segnaposto che riceve: deve tornare un nome."""

    def scrivi(self, sistema, richiesta) -> Risposta:
        segno = re.search(r"\[SOGGETTO_\d+\]", richiesta).group(0)
        testo = f"Per {segno}, opponente. {{fonte: modello}}"
        return Risposta(testo, "registratore", 100, 10)


def test_un_profilo_non_confermato_non_scrive(app, impostazioni, civile):
    with app.state.session_factory() as db:
        sarti = utente(db, "sarti")
        decidi_formule(db, sarti, [1], "tolta")  # ora non è più confermato
        with pytest.raises(BozzaRifiutata, match="profilo, confermato"):
            prepara_bozza(db, sarti, civile, "opposizione", ModelloFinto())


def test_partono_solo_le_formule_tenute_e_le_scadenze_confermate(app, civile):
    modello = ModelloCheRicorda()
    with app.state.session_factory() as db:
        scadenza = db.scalar(select(Scadenza))
        scadenza.confermata_da_id = None
        db.commit()
        prepara_bozza(db, utente(db, "sarti"), civile, "opposizione", modello)
    (inviato,) = modello.ricevuti
    assert "transazione commerciale tra imprese" not in inviato
    assert "SCADENZE:\nnessuna" in inviato


def test_nella_bozza_tornano_i_nomi_veri(app, civile):
    with app.state.session_factory() as db:
        esito = prepara_bozza(
            db, utente(db, "sarti"), civile, "opposizione", CitaIlCliente()
        )
    (paragrafo,) = esito.paragrafi
    assert paragrafo.testo.startswith("Per Ristorazione Collinare S.r.l.")
    assert paragrafo.problemi == []


def test_i_nomi_tornano_e_i_segnaposto_dello_stile_si_segnalano(app, civile):
    with app.state.session_factory() as db:
        esito = prepara_bozza(
            db, utente(db, "sarti"), civile, "opposizione", ModelloFinto()
        )
    testo = " ".join(p.testo for p in esito.paragrafi)
    assert "TRIBUNALE ORDINARIO DI MODENA" in testo
    rimasti = [m for p in esito.paragrafi for m in p.problemi]
    assert any(m.startswith("segnaposto rimasto: [") for m in rimasti)
    assert esito.stile == "Elena Sarti"


def test_il_promemoria_dice_entro_quando(app, civile):
    with app.state.session_factory() as db:
        esito = prepara_bozza(
            db, utente(db, "irene"), civile, "promemoria", ModelloFinto()
        )
    testo = " ".join(p.testo for p in esito.paragrafi)
    assert "lunedì 2 novembre 2026" in testo
    assert not any("entro quando" in n for n in esito.note)


def test_senza_scadenza_confermata_il_promemoria_lo_dice(app, impostazioni):
    fascicolo = con_fascicoli(app)["2026-071"]
    with app.state.session_factory() as db:
        esito = prepara_bozza(
            db, utente(db, "irene"), fascicolo, "promemoria", ModelloFinto()
        )
    assert any(n.startswith("Nessuna scadenza confermata") for n in esito.note)


def test_un_nome_fuori_elenco_blocca_la_bozza(app, impostazioni):
    # Il decreto va a capo dopo «Termocucine» (lezione 18): senza la
    # controparte in elenco, la bozza non parte.
    fascicolo = con_fascicoli(app)["2026-071"]
    with app.state.session_factory() as db:
        irene = utente(db, "irene")
        decreto = CIVILE / "decreto.pdf"
        app_documenti.carica(
            db,
            impostazioni,
            irene,
            db.get(Fascicolo, fascicolo.id),
            decreto.name,
            decreto.read_bytes(),
        )
        modello = ModelloCheRicorda()
        with pytest.raises(InvioBloccato):
            prepara_bozza(db, irene, fascicolo, "promemoria", modello)
    assert modello.ricevuti == []


def test_chi_non_puo_e_che_cosa_non_si_fa(app, civile):
    with app.state.session_factory() as db:
        with pytest.raises(RichiestaNonAmmessa):
            prepara_bozza(
                db, utente(db, "rosa"), civile, "promemoria", ModelloFinto()
            )
        with pytest.raises(TipoDiBozza):
            prepara_bozza(
                db, utente(db, "sarti"), civile, "appello", ModelloFinto()
            )


def test_in_word_secondo_il_dm_110_2023(app, civile, tmp_path):
    with app.state.session_factory() as db:
        esito = prepara_bozza(
            db, utente(db, "sarti"), civile, "opposizione", ModelloFinto()
        )
        percorso = in_word(db, esito.bozza, tmp_path / "bozza.docx")
    documento = Document(str(percorso))
    sezione = documento.sections[0]
    # Word conta in ventesimi di punto: 2,5 cm tornano con l'arrotondamento.
    for margine in (sezione.left_margin, sezione.top_margin):
        assert abs(margine - Cm(2.5)) < Cm(0.01)
    normale = documento.styles["Normal"]
    assert normale.font.size == Pt(12)
    assert normale.paragraph_format.line_spacing == 1.5
    gialli = [
        r.text
        for p in documento.paragraphs
        for r in p.runs
        if r.font.highlight_color == WD_COLOR_INDEX.YELLOW
    ]
    assert any(t.startswith("[DA COMPLETARE") for t in gialli)
    with zipfile.ZipFile(percorso) as archivio:
        commenti = archivio.read("word/comments.xml").decode("utf-8")
    assert "Provenienza: stile" in commenti
    riquadro = documento.tables[0].cell(0, 0).text
    assert "importo che non è nei documenti: 5.000,00" in riquadro


def test_la_pagina_delle_bozze_e_il_file(app, browser, civile):
    sarti = browser("sarti")
    indirizzo = f"/fascicoli/{civile.id}/bozze"
    assert invia(sarti, indirizzo, {"tipo": "promemoria"}).status_code == 303
    pagina = sarti.get(indirizzo).text
    assert "lunedì 2 novembre 2026" in pagina
    with app.state.session_factory() as db:
        bozza = db.scalar(select(Bozza))
    file = sarti.get(f"{indirizzo}/{bozza.id}/word")
    assert file.status_code == 200 and file.content.startswith(b"PK")


def test_il_comando_mostra_provenienza_e_problemi(app, civile, capsys):
    with app.state.session_factory() as db:
        settings = app.state.settings
        assert (
            manage.bozza_dalla_riga(
                db,
                settings,
                Namespace(
                    codice="2026-071",
                    tipo="opposizione",
                    da="sarti",
                    word=None,
                    verbale=None,
                ),
            )
            == 0
        )
        voci = db.scalar(select(func.count()).select_from(UsoAI))
    uscita = capsys.readouterr().out
    assert "--- la bozza, con la provenienza di ogni paragrafo ---" in uscita
    assert "! importo che non è nei documenti: 5.000,00" in uscita
    assert voci == 2  # la scheda di prova e la bozza


def test_la_migrazione_10_aggiunge_le_bozze(impostazioni):
    tutti = dict(PASSI)
    engine = make_engine(impostazioni.database_url)
    try:
        PASSI.clear()
        PASSI.update({n: tutti[n] for n in range(1, 10)})  # la versione 0.9
        migra(engine)
        PASSI[10] = tutti[10]  # arriva la versione 1.0
        assert migra(engine) == [10]
        with engine.connect() as conn:
            assert inspect(conn).has_table("bozze")
    finally:
        PASSI.clear()
        PASSI.update(tutti)
        engine.dispose()


def test_le_note_restano_con_la_bozza(app, civile):
    with app.state.session_factory() as db:
        prepara_bozza(
            db, utente(db, "sarti"), civile, "opposizione", ModelloFinto()
        )
        bozza = db.scalar(select(Bozza))
    assert any(n.startswith("Da verificare") for n in json.loads(bozza.note))


# [libro:prova-approvazione]
def test_si_approva_il_file_corretto_non_la_bozza(app, civile, tmp_path):
    with app.state.session_factory() as db:
        sarti, fascicolo = utente(db, "sarti"), db.get(Fascicolo, civile.id)
        bozza = prepara_bozza(
            db, sarti, fascicolo, "promemoria", ModelloFinto()
        ).bozza
        file = in_word(db, bozza, tmp_path / "bozza.docx")
        with pytest.raises(BozzaRifiutata, match="il riquadro di Minuta"):
            approva_bozza(db, sarti, fascicolo, bozza, file)
        corretto = Document()  # il testo che l'avvocato ha riscritto
        corretto.add_paragraph("Gentile cliente, ecco che cosa è arrivato.")
        corretto.save(str(tmp_path / "corretto.docx"))
        approva_bozza(db, sarti, fascicolo, bozza, tmp_path / "corretto.docx")
        assert bozza.approvata_da_id == sarti.id and bozza.forzatura is None
        assert bozza.testo_approvato.startswith("Gentile cliente")


def test_la_forzatura_resta_nel_registro(app, civile):
    with app.state.session_factory() as db:
        sarti, fascicolo = utente(db, "sarti"), db.get(Fascicolo, civile.id)
        bozza = prepara_bozza(
            db, sarti, fascicolo, "promemoria", ModelloFinto()
        ).bozza
        with pytest.raises(BozzaRifiutata, match="dato da completare"):
            approva_bozza(db, sarti, fascicolo, bozza)
        residui = approva_bozza(db, sarti, fascicolo, bozza, forza=True)
        evento = db.scalar(
            select(Evento).where(Evento.azione == "approvazione")
        )
    assert residui == ["1 dato da completare"]
    assert json.loads(evento.dettaglio)["forzatura"] == residui


# [/libro:prova-approvazione]


def test_approva_solo_un_avvocato_e_una_volta(app, civile):
    with app.state.session_factory() as db:
        irene, fascicolo = utente(db, "irene"), db.get(Fascicolo, civile.id)
        bozza = prepara_bozza(
            db, irene, fascicolo, "promemoria", ModelloFinto()
        ).bozza
        with pytest.raises(BozzaRifiutata, match="un avvocato"):
            approva_bozza(db, irene, fascicolo, bozza, forza=True)
        approva_bozza(db, utente(db, "sarti"), fascicolo, bozza, forza=True)
        with pytest.raises(BozzaRifiutata, match="già approvata"):
            approva_bozza(db, utente(db, "sarti"), fascicolo, bozza)


@pytest.fixture
def massimario_di_prova(monkeypatch, tmp_path):
    """Una copia del massimario: le prove non scrivono in quello vero."""
    copia = tmp_path / "massimario.json"
    copia.write_text(app_bozze.MASSIMARIO.read_text("utf-8"), "utf-8")
    monkeypatch.setattr(app_bozze, "MASSIMARIO", copia)
    return copia


def test_le_citazioni_della_bozza_si_controllano(app, civile):
    with app.state.session_factory() as db:
        esito = prepara_bozza(
            db, utente(db, "sarti"), civile, "opposizione", ModelloFinto()
        )
    problemi = [m for p in esito.paragrafi for m in p.problemi]
    assert any(
        m.startswith("sentenza che il massimario non ha") for m in problemi
    )
    assert any(m.startswith("parla di assegno o cambiale") for m in problemi)
    assert "Da leggere sulla fonte: c.p.c. art. 645" in esito.note


def test_una_norma_letta_entra_nel_massimario(
    app, civile, massimario_di_prova
):
    with app.state.session_factory() as db:
        sarti = utente(db, "sarti")
        bozza = prepara_bozza(
            db, sarti, civile, "opposizione", ModelloFinto()
        ).bozza
        prima = dict((c, s) for _n, c, s in citazioni_della_bozza(bozza))
        voce = norma_letta(db, sarti, "c.p.c. art. 645")
        dopo = dict((c, s) for _n, c, s in citazioni_della_bozza(bozza))
        with pytest.raises(BozzaRifiutata, match="già nel massimario"):
            norma_letta(db, sarti, "c.p.c. art. 645")
        with pytest.raises(BozzaRifiutata, match="un avvocato"):
            norma_letta(db, utente(db, "irene"), "c.p.c. art. 646")
        with pytest.raises(BozzaRifiutata, match="non so dove leggerla"):
            norma_letta(db, sarti, "art. 645")
    assert prima["c.p.c. art. 645"] == "da leggere sulla fonte"
    assert dopo["c.p.c. art. 645"].startswith("letta da Elena Sarti il ")
    assert voce["indirizzo"].endswith("1443:1~art645")
    assert (
        json.loads(massimario_di_prova.read_text("utf-8"))["c.p.c. art. 645"][
            "letta_da"
        ]
        == "Elena Sarti"
    )


def test_una_sentenza_non_letta_non_entra_nell_atto(app, civile, tmp_path):
    with app.state.session_factory() as db:
        sarti, fascicolo = utente(db, "sarti"), db.get(Fascicolo, civile.id)
        bozza = prepara_bozza(
            db, sarti, fascicolo, "opposizione", ModelloFinto()
        ).bozza
        corretto = Document()
        corretto.add_paragraph("Lo dice Cass. civ., 15 maggio 2023, n. 99999.")
        corretto.save(str(tmp_path / "corretto.docx"))
        with pytest.raises(BozzaRifiutata, match="sentenza che il massimario"):
            approva_bozza(
                db, sarti, fascicolo, bozza, tmp_path / "corretto.docx"
            )
