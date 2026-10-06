import json
import shutil

import pytest
from conftest import RADICE
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from minuta import __main__ as comandi
from minuta import archive, draft, learning, word
from minuta.model import ModelloFinto


@pytest.fixture
def bozza_word(db, fascicolo, config, profilo, massimario, tmp_path):
    """Una bozza del modello finto, in Markdown e in Word."""
    bozza = draft.prepara(db, fascicolo, config, profilo, massimario, ModelloFinto(), tmp_path)
    markdown = draft.in_markdown(bozza, fascicolo)
    (tmp_path / "bozza.md").write_text(markdown, "utf-8")
    return markdown, word.scrivi(markdown, tmp_path / "bozza.docx")


def con_revisione(percorso, vecchio, nuovo, uscita):
    """Come l'avvocato con le revisioni attive: «vecchio» cancellato, «nuovo» inserito."""
    documento = Document(str(percorso))
    paragrafo = next(p for p in documento.paragraphs if vecchio in p.text)
    prima, dopo = paragrafo.text.split(vecchio, 1)
    for run in paragrafo.runs:
        run._r.getparent().remove(run._r)
    paragrafo.add_run(prima)
    for tag, testo, nome in (("w:del", vecchio, "w:delText"), ("w:ins", nuovo, "w:t")):
        revisione = OxmlElement(tag)
        revisione.set(qn("w:id"), str(900 + len(tag)))
        revisione.set(qn("w:author"), "Avv. Elena Sarti")
        run, elemento = OxmlElement("w:r"), OxmlElement(nome)
        elemento.text = testo
        elemento.set(qn("xml:space"), "preserve")
        run.append(elemento)
        revisione.append(run)
        paragrafo._p.append(revisione)
    paragrafo.add_run(dopo)
    documento.save(str(uscita))
    return uscita


def ripulito(percorso, sostituzioni, uscita):
    """Come l'avvocato prima del deposito: corregge, toglie riquadro e commenti di Minuta."""
    documento = Document(str(percorso))
    corpo = documento.element.body
    for tabella in documento.tables:
        if tabella.cell(0, 0).paragraphs[0].style.name == word.STILE_NOTA:
            corpo.remove(tabella._tbl)
    for paragrafo in documento.paragraphs:
        for run in paragrafo.runs:
            for vecchio, nuovo in sostituzioni:
                run.text = run.text.replace(vecchio, nuovo)
    for nome in ("w:commentRangeStart", "w:commentRangeEnd"):
        for elemento in list(corpo.iter(qn(nome))):
            elemento.getparent().remove(elemento)
    for riferimento in list(corpo.iter(qn("w:commentReference"))):
        run = riferimento.getparent()
        run.getparent().remove(run)
    commenti = documento.part.part_related_by(RT.COMMENTS).element
    for commento in list(commenti):
        commenti.remove(commento)
    documento.save(str(uscita))
    return uscita


def test_senza_correzioni_il_giro_in_word_non_cambia_nulla(bozza_word):
    markdown, percorso = bozza_word
    testo = "\n\n".join(word.leggi(percorso))
    assert learning.confronta(learning.testo_della_bozza(markdown), testo) == []


def test_l_impaginazione_segue_il_dm_110_2023(bozza_word):
    documento = Document(str(bozza_word[1]))
    sezione = documento.sections[0]
    assert (round(sezione.page_width.cm, 1), round(sezione.page_height.cm, 1)) == (21.0, 29.7)
    for margine in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        assert round(getattr(sezione, margine).cm, 2) == 2.5
    normale = documento.styles["Normal"]
    assert normale.font.size.pt == 12 and normale.paragraph_format.line_spacing == 1.5
    assert len(normale.element.findall(f".//{qn('w:lang')}")) == 1


def test_la_provenienza_sta_nei_commenti_a_margine(bozza_word):
    markdown, percorso = bozza_word
    commenti = list(Document(str(percorso)).comments)
    fonti = [p["fonte"] for p in word.da_markdown(markdown)["paragrafi"] if p["fonte"]]
    assert [c.author for c in commenti] == ["Minuta"] * len(fonti)
    assert [c.text for c in commenti] == [f"Provenienza: {f}" for f in fonti]


def test_i_dati_da_completare_sono_evidenziati(bozza_word):
    runs = [r for p in Document(str(bozza_word[1])).paragraphs for r in p.runs
            if r.text.startswith("[DA COMPLETARE")]
    assert runs and all(r.font.highlight_color == WD_COLOR_INDEX.YELLOW for r in runs)


def test_le_revisioni_si_leggono_come_accettate(bozza_word, tmp_path):
    rivisto = con_revisione(bozza_word[1], "la somma di euro", "la sorte capitale di euro",
                            tmp_path / "rivisto.docx")
    testo = "\n".join(word.leggi(rivisto))
    assert "la sorte capitale di euro" in testo and "la somma di euro" not in testo
    assert "2 revisioni non accettate" in word.residui(rivisto)


def test_un_a_capo_cancellato_unisce_i_paragrafi(tmp_path):
    documento = Document()
    primo = documento.add_paragraph("Modena, ")
    proprieta = OxmlElement("w:rPr")
    cancellato = OxmlElement("w:del")
    cancellato.set(qn("w:id"), "1")
    cancellato.set(qn("w:author"), "Avv. Elena Sarti")
    proprieta.append(cancellato)
    primo._p.get_or_add_pPr().append(proprieta)
    documento.add_paragraph("20 ottobre 2026")
    tabella = documento.add_table(rows=1, cols=1)
    tabella.cell(0, 0).text = "fattura n. 455/2025"
    documento.save(str(tmp_path / "a-capo.docx"))
    assert word.leggi(tmp_path / "a-capo.docx") == ["Modena, 20 ottobre 2026", "fattura n. 455/2025"]


def test_cio_che_resta_di_minuta_si_vede(bozza_word, tmp_path):
    residui = word.residui(bozza_word[1])
    assert residui[0].startswith("il riquadro di Minuta (")
    assert any(r.endswith("commenti di Minuta") for r in residui)
    assert "1 dato da completare" in residui
    firmato = ripulito(bozza_word[1], [("[DA COMPLETARE: luogo e data]", "Modena, 20 ottobre 2026,")],
                       tmp_path / "firmato.docx")
    assert word.residui(firmato) == []


def test_la_carta_intestata_dello_studio(tmp_path):
    carta = Document()
    carta.sections[0].header.paragraphs[0].text = "Studio Meridiana · Via delle Meridiane 12 · Modena"
    carta.add_paragraph("testo del modello, da non tenere")
    carta.save(str(tmp_path / "carta.docx"))
    markdown = ("> **BOZZA** · fascicolo 1 · prova\n\nTRIBUNALE DI MODENA\n\n"
                "<sub>provenienza: fascicolo</sub>\n\n---\n\n## Note per l'avvocato\n\n- Nessun avviso.\n")
    uscita = word.scrivi(markdown, tmp_path / "bozza.docx", carta=tmp_path / "carta.docx")
    documento = Document(str(uscita))
    assert documento.sections[0].header.paragraphs[0].text.startswith("Studio Meridiana")
    assert "testo del modello" not in "\n".join(p.text for p in documento.paragraphs)
    assert word.leggi(uscita) == ["TRIBUNALE DI MODENA"]


def test_il_giro_completo_dai_comandi(fascicolo, config, tmp_path, monkeypatch, capsys):
    (tmp_path / "config").mkdir()
    for nome in ("studio.json", "profilo.json", "massimario.json"):
        shutil.copy(RADICE / "config" / nome, tmp_path / "config" / nome)
    archive.importa(RADICE / "archivio/pdf", archive.apri(tmp_path / "minuta.db"), config)
    (tmp_path / "fascicolo.json").write_text(json.dumps(fascicolo), "utf-8")
    monkeypatch.setattr(comandi, "RADICE", tmp_path)

    assert comandi.main(["bozza", str(tmp_path / "fascicolo.json"), "--modello", "finto",
                         "--uscita", str(tmp_path / "bozze")]) == 0
    in_word = next((tmp_path / "bozze").glob("*.docx"))
    assert in_word.with_suffix(".md").exists()
    approva = ["approva", str(in_word), "--avvocato", "sarti",
               "--fascicolo", str(tmp_path / "fascicolo.json"), "--finale"]

    # La bozza non corretta non passa per un atto firmato.
    with pytest.raises(SystemExit):
        comandi.main([*approva, str(in_word)])
    errore = capsys.readouterr().err
    assert "l'atto sembra ancora una bozza (il riquadro di Minuta" in errore
    assert "1 dato da completare" in errore

    firmato = ripulito(in_word, [("la somma di euro", "la sorte capitale di euro"),
                                 ("[DA COMPLETARE: luogo e data]", "Modena, 20 ottobre 2026,")],
                       tmp_path / "firmato.docx")
    assert comandi.main([*approva, str(firmato)]) == 0
    uscita = capsys.readouterr().out
    assert ("approvata da sarti: 1 completamento, 0 modifiche di dati, 1 correzione di stile, "
            "0 riscritture") in uscita
    assert "stile: «somma» → «sorte capitale»" in uscita

    # Con --comunque si registra, e il registro dice che cosa restava.
    assert comandi.main([*approva, str(in_word), "--comunque"]) == 0
    assert "ATTENZIONE, nell'atto resta: 1 dato da completare" in capsys.readouterr().out
    eventi = [json.loads(r) for r in (tmp_path / "registro/uso-ai.jsonl").read_text().splitlines()]
    approvazioni = [e for e in eventi if e.get("evento") == "approvazione"]
    assert approvazioni[0]["residui"] == []
    assert "1 dato da completare" in approvazioni[1]["residui"]


def test_una_bozza_gia_fatta_si_porta_in_word(bozza_word, config, tmp_path, monkeypatch, capsys):
    markdown, _ = bozza_word
    (tmp_path / "config").mkdir(exist_ok=True)
    shutil.copy(RADICE / "config/studio.json", tmp_path / "config/studio.json")
    monkeypatch.setattr(comandi, "RADICE", tmp_path)
    (tmp_path / "vecchia.md").write_text(markdown, "utf-8")
    assert comandi.main(["word", str(tmp_path / "vecchia.md")]) == 0
    assert "bozza in Word:" in capsys.readouterr().out
    assert word.leggi(tmp_path / "vecchia.docx") == word.leggi(bozza_word[1])
